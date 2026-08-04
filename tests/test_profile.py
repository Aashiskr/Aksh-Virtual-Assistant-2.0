import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from docx import Document

from backend.brain.groq import GroqBrain
from backend.config import AkshSettings
from backend.profile import UserProfileStore


class UserProfileStoreTests(unittest.TestCase):
    def test_text_cv_is_extracted_and_persisted(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "resume.txt"
            source.write_text(
                "Demo User\nPython developer\nBuilt an AI assistant.",
                encoding="utf-8",
            )
            store = UserProfileStore(root / "data")
            imported = store.import_file(source)
            self.assertEqual(imported["source_name"], "resume.txt")
            self.assertIn("Python developer", store.context())
            self.assertNotIn(str(source), store.context())

    def test_docx_cv_tables_and_paragraphs_are_extracted(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "resume.docx"
            document = Document()
            document.add_paragraph(
                "Demo User — backend software developer"
            )
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Skill"
            table.cell(0, 1).text = "Python"
            document.save(source)
            store = UserProfileStore(root / "data")
            store.import_file(source)
            context = store.context()
            self.assertIn("Demo User", context)
            self.assertIn("Skill | Python", context)

    def test_profile_can_be_removed(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "profile.md"
            source.write_text(
                "Software engineer with backend and automation experience.",
                encoding="utf-8",
            )
            store = UserProfileStore(root / "data")
            store.import_file(source)
            self.assertTrue(store.remove())
            self.assertEqual(store.load(), {})

    def test_normal_aksh_brain_receives_owner_profile_context(self):
        brain = GroqBrain(AkshSettings())
        brain.profile = Mock()
        brain.profile.context.return_value = "Owner knows Python and FastAPI."
        self.assertIn(
            "Owner knows Python and FastAPI.",
            brain._system_prompt(),
        )


if __name__ == "__main__":
    unittest.main()
