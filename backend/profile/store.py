from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".rst",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
}
MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_PROFILE_CHARACTERS = 60000


class UserProfileStore:
    """Stores owner-selected CV text locally for personalized Aksh context."""

    def __init__(self, data_dir: Path):
        self.path = data_dir / "user_profile.json"

    def import_file(self, source: Path) -> dict[str, Any]:
        source = source.expanduser().resolve()
        if not source.is_file():
            raise ValueError("Selected CV file nahi mili.")
        if source.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError("PDF, DOCX, TXT, Markdown ya image CV select karein.")
        if source.stat().st_size > MAX_FILE_BYTES:
            raise ValueError("CV file 12 MB se chhoti honi chahiye.")
        text = _extract_text(source)
        text = _clean_text(text)[:MAX_PROFILE_CHARACTERS]
        if len(text) < 30:
            raise ValueError("CV se readable text extract nahi hua.")
        value = {
            "source_name": source.name,
            "imported_at": datetime.now().astimezone().isoformat(),
            "characters": len(text),
            "text": text,
        }
        self._save(value)
        return value

    def load(self) -> dict[str, Any]:
        try:
            value = json.loads(self.path.read_text(encoding="utf-8"))
            return value if isinstance(value, dict) else {}
        except (OSError, json.JSONDecodeError):
            return {}

    def context(self, maximum_characters: int = 12000) -> str:
        value = self.load()
        text = str(value.get("text", "")).strip()
        if not text:
            return ""
        name = str(value.get("source_name", "Owner profile"))
        return f"Source: {name}\n{text[:maximum_characters]}"

    def remove(self) -> bool:
        if not self.path.exists():
            return False
        self.path.unlink()
        return True

    def _save(self, value: dict[str, Any]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.path.with_suffix(".tmp")
        temporary.write_text(
            json.dumps(value, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        temporary.replace(self.path)


def _extract_text(source: Path) -> str:
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(source))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(str(source))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix in {".txt", ".md", ".rst"}:
        return source.read_text(encoding="utf-8", errors="replace")
    from PIL import Image
    import pytesseract

    with Image.open(source) as image:
        return str(pytesseract.image_to_string(image))


def _clean_text(text: str) -> str:
    clean = str(text).replace("\x00", " ")
    clean = re.sub(r"[ \t]+", " ", clean)
    clean = re.sub(r"\n{3,}", "\n\n", clean)
    return clean.strip()
